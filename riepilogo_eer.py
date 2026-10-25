"""
RENTRI Analyzer

Copyright (C) 2026 Gianluca Massei

Autore:
Gianluca Massei

Questo software è distribuito con licenza MIT.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND.
"""
from pathlib import Path

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.styles import Font

from fifo import calcola_fifo


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="4472C4"
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True
)


def genera_riepilogo_eer(
    df,
    operazioni_finali,
    output_file="riepilogo_eer.xlsx"
):

    wb = Workbook()

    ws = wb.active

    ws.title = "RIEPILOGO_EER"

    ws.append([
        "EER",
        "Giacenza Max kg",
        "Data Giacenza Max",
        "Quantità Annua kg",
        "Giorni Max Deposito",
        "Movimenti",
        "Lotti Residui"
    ])

    for cella in ws[1]:

        cella.fill = HEADER_FILL
        cella.font = HEADER_FONT

    eers = sorted(
        df["Codice EER"]
        .dropna()
        .astype(str)
        .unique()
    )

    for eer in eers:

        tmp = df[
            df["Codice EER"]
            .astype(str)
            == eer
        ].copy()

        tmp = tmp.sort_values(
            ["Data registrazione","Progressivo"]
        )

        tmp["Delta"] = tmp.apply(
            lambda r:
            r["Quantita"]
            if r["Tipo operazione"] == "Carico"
            else -r["Quantita"],
            axis=1
        )

        tmp["Giacenza"] = (
            tmp["Delta"].cumsum()
        )

        giacenza_max = (
            tmp["Giacenza"].max()
        )

        idx_max = (
            tmp["Giacenza"].idxmax()
        )

        data_max = (
            tmp.loc[
                idx_max,
                "Data registrazione"
            ]
        )

        quantita_annua = (
            tmp[
                tmp["Tipo operazione"]
                == "Carico"
            ]["Quantita"]
            .sum()
        )

        fifo_input = tmp[
            (
                tmp["Tipo operazione"]
                == "Carico"
            )
            |
            (
                (
                    tmp["Tipo operazione"]
                    == "Scarico"
                )
                &
                (
                    tmp["Destinato a"]
                    .fillna("")
                    .isin(
                        operazioni_finali
                    )
                )
            )
        ].copy()

        fifo_df, residui_df = (
            calcola_fifo(
                fifo_input
            )
        )

        giorni_max = 0

        if not fifo_df.empty:

            giorni_max = (
                fifo_df[
                    "GiorniDeposito"
                ].max()
            )

        ws.append([

            eer,

            float(giacenza_max),

            data_max,

            float(quantita_annua),

            int(giorni_max),

            len(tmp),

            len(residui_df)

        ])

    wb.save(
        output_file
    )

    return Path(output_file)