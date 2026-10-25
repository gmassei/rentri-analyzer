"""
RENTRI Analyzer

Copyright (C) 2026 Gianluca Massei

Autore:
Gianluca Massei

Questo software è distribuito con licenza MIT.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="4472C4"
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True
)

ROSSO = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)

GIALLO = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC"
)


def formatta_header(ws):

    for cella in ws[1]:
        cella.fill = HEADER_FILL
        cella.font = HEADER_FONT


def genera_report_excel(
    eer,
    movimenti_df,
    fifo_df,
    residui_df,
    operazioni_finali,
    limite_istantaneo=None,
    limite_annuo=None,
    limite_giorni=None,
    output_file="report.xlsx"
):

    wb = Workbook()

    # =====================================================
    # RIEPILOGO
    # =====================================================

    ws_riep = wb.active
    ws_riep.title = "RIEPILOGO"

    ws_riep.append([
        "Parametro",
        "Valore"
    ])

    formatta_header(ws_riep)

    giacenza_max = float(
        movimenti_df["Giacenza"].max()
    )

    carichi_annui = (
        movimenti_df[
            movimenti_df["Tipo operazione"] == "Carico"
        ]
        .groupby(
            movimenti_df["Data registrazione"].dt.year
        )["Quantita"]
        .sum()
    )

    quantita_annua_massima = (
        float(carichi_annui.max())
        if len(carichi_annui)
        else 0
    )

    giorni_max = 0

    if not fifo_df.empty:

        giorni_max = int(
            fifo_df["GiorniDeposito"].max()
        )

    ws_riep.append(["EER", eer])
    ws_riep.append(["Giacenza massima", giacenza_max])
    ws_riep.append(["Quantitativo annuo massimo", quantita_annua_massima])
    ws_riep.append(["Giorni massimi deposito", giorni_max])

    ws_riep.append([])
    ws_riep.append(["Limite istantaneo", limite_istantaneo])
    ws_riep.append(["Limite annuo", limite_annuo])
    ws_riep.append(["Limite deposito giorni", limite_giorni])

    ws_riep.append([])
    ws_riep.append([
        "Operazioni finali considerate"
    ])

    for op in operazioni_finali:

        ws_riep.append([op])

    # =====================================================
    # MOVIMENTI
    # =====================================================

    ws_mov = wb.create_sheet(
        "MOVIMENTI"
    )

    ws_mov.append([
        "Data",
        "Progressivo",
        "Tipo operazione",
        "Destinato a",
        "Quantita",
        "Giacenza"
    ])

    formatta_header(ws_mov)

    for _, r in movimenti_df.iterrows():

        ws_mov.append([
            r["Data registrazione"],
            r["Progressivo"],
            r["Tipo operazione"],
            r["Destinato a"],
            r["Quantita"],
            r["Giacenza"]
        ])

        if (
            limite_istantaneo
            and
            r["Giacenza"] > limite_istantaneo
        ):

            for c in ws_mov[ws_mov.max_row]:
                c.fill = ROSSO

    # =====================================================
    # FIFO
    # =====================================================

    ws_fifo = wb.create_sheet(
        "FIFO"
    )

    ws_fifo.append([
        "DataCarico",
        "ProgressivoCarico",
        "DataScarico",
        "ProgressivoScarico",
        "Quantita",
        "GiorniDeposito",
        "Superamento"
    ])

    formatta_header(ws_fifo)

    if not fifo_df.empty:

        for _, r in fifo_df.iterrows():

            superamento = False

            if (
                limite_giorni
                and
                r["GiorniDeposito"] > limite_giorni
            ):
                superamento = True

            ws_fifo.append([
                r["DataCarico"],
                r["ProgressivoCarico"],
                r["DataScarico"],
                r["ProgressivoScarico"],
                r["Quantita"],
                r["GiorniDeposito"],
                superamento
            ])

            if superamento:

                for c in ws_fifo[
                    ws_fifo.max_row
                ]:
                    c.fill = ROSSO

    # =====================================================
    # RESIDUI
    # =====================================================

    ws_res = wb.create_sheet(
        "RESIDUI"
    )

    ws_res.append([
        "DataCarico",
        "ProgressivoCarico",
        "QuantitaResidua"
    ])

    formatta_header(ws_res)

    if not residui_df.empty:

        for _, r in residui_df.iterrows():

            ws_res.append([
                r["DataCarico"],
                r["ProgressivoCarico"],
                r["QuantitaResidua"]
            ])

    # =====================================================
    # ANOMALIE
    # =====================================================

    ws_an = wb.create_sheet(
        "ANOMALIE"
    )

    ws_an.append([
        "Tipo",
        "Data",
        "Valore",
        "Limite"
    ])

    formatta_header(ws_an)

    # anomalie quantità

    if limite_istantaneo:

        sup_qta = movimenti_df[
            movimenti_df["Giacenza"]
            > limite_istantaneo
        ]

        for _, r in sup_qta.iterrows():

            ws_an.append([
                "Giacenza",
                r["Data registrazione"],
                r["Giacenza"],
                limite_istantaneo
            ])

            for c in ws_an[
                ws_an.max_row
            ]:
                c.fill = ROSSO

    # anomalie tempo

    if (
        not fifo_df.empty
        and
        limite_giorni
    ):

        sup_tmp = fifo_df[
            fifo_df["GiorniDeposito"]
            > limite_giorni
        ]

        for _, r in sup_tmp.iterrows():

            ws_an.append([
                "Tempo deposito",
                r["DataScarico"],
                r["GiorniDeposito"],
                limite_giorni
            ])

            for c in ws_an[
                ws_an.max_row
            ]:
                c.fill = ROSSO

    wb.save(output_file)

    return Path(output_file)